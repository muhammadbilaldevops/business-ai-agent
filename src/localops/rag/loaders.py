"""Bounded document extraction. Archives are inspected before Office parsing."""

import csv
import io
import json
import zipfile
from pathlib import Path

from fastapi import HTTPException

from localops.security import limit_bytes, safe_filename


def extract(filename: str, content: bytes, maximum: int = 1_000_000) -> list[dict]:
    safe_filename(filename)
    limit_bytes(content, 10 * 1024 * 1024)
    extension = Path(filename).suffix.lower()
    if extension not in {".md", ".txt", ".pdf", ".docx", ".csv", ".json", ".xlsx"}:
        raise HTTPException(415, "Supported formats: PDF, DOCX, Markdown, TXT, CSV, JSON, XLSX")
    try:
        if extension in {".docx", ".xlsx"}:
            with zipfile.ZipFile(io.BytesIO(content)) as z:
                entries = z.infolist()
                if len(entries) > 2000 or sum(i.file_size for i in entries) > 20_000_000:
                    raise HTTPException(413, "Expanded Office document is too large")
        if extension == ".pdf":
            from pypdf import PdfReader

            reader = PdfReader(io.BytesIO(content))
            if reader.is_encrypted or len(reader.pages) > 200:
                raise HTTPException(400, "Use an unencrypted PDF with at most 200 pages")
            pages = [
                {"text": page.extract_text() or "", "page": i + 1}
                for i, page in enumerate(reader.pages)
            ]
        elif extension == ".docx":
            from docx import Document

            doc = Document(io.BytesIO(content))
            parts = [p.text for p in doc.paragraphs] + [
                " | ".join(c.text for c in row.cells) for table in doc.tables for row in table.rows
            ]
            pages = [{"text": "\n".join(parts), "page": None}]
        elif extension == ".xlsx":
            columns, rows = load_table(filename, content)
            pages = [
                {
                    "text": "\n".join(
                        [",".join(columns)]
                        + [",".join(str(r.get(c, "")) for c in columns) for r in rows]
                    ),
                    "page": None,
                }
            ]
        else:
            text = content.decode("utf-8-sig")
            if extension == ".json":
                text = json.dumps(json.loads(text), ensure_ascii=False, indent=2)
            pages = [{"text": text, "page": None}]
        if sum(len(p["text"]) for p in pages) > maximum:
            raise HTTPException(413, "Extracted text exceeds the workspace limit")
        if not any(p["text"].strip() for p in pages):
            raise HTTPException(422, "No text found. Scanned PDFs need OCR before uploading.")
        return pages
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(
            422, "The document could not be read; check its format and encoding"
        ) from exc


def load_table(filename: str, content: bytes) -> tuple[list[str], list[dict]]:
    limit_bytes(content, 10 * 1024 * 1024)
    try:
        if filename.lower().endswith(".xlsx"):
            from openpyxl import load_workbook

            with zipfile.ZipFile(io.BytesIO(content)) as z:
                if sum(i.file_size for i in z.infolist()) > 20_000_000:
                    raise HTTPException(413, "Expanded workbook is too large")
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            try:
                source = workbook.active.iter_rows(values_only=True)
                values = []
                for i, row in enumerate(source):
                    if i > 10000:
                        raise HTTPException(413, "Use at most 10,000 data rows")
                    values.append(list(row))
            finally:
                workbook.close()
        elif filename.lower().endswith(".csv"):
            values = []
            for i, row in enumerate(csv.reader(io.StringIO(content.decode("utf-8-sig")))):
                if i > 10000:
                    raise HTTPException(413, "Use at most 10,000 data rows")
                values.append(row)
        else:
            raise HTTPException(415, "Upload CSV or XLSX")
        if len(values) < 2:
            raise HTTPException(422, "A dataset needs column headers and at least one row")
        columns = [str(v or "").strip() for v in values[0]]
        if (
            not 1 <= len(columns) <= 100
            or any(not c or len(c) > 100 for c in columns)
            or len({c.lower() for c in columns}) != len(columns)
        ):
            raise HTTPException(422, "Use 1–100 distinct, non-empty column names")
        if any(len(row) != len(columns) for row in values[1:]):
            raise HTTPException(422, "Every row must have the same number of columns")
        rows = [
            {c: (str(v) if v is not None else "") for c, v in zip(columns, row)}
            for row in values[1:]
        ]
        if sum(sum(len(v) for v in row.values()) for row in rows) > 1_000_000:
            raise HTTPException(413, "Dataset cells exceed the text limit")
        return columns, rows
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(422, "Could not parse dataset") from exc
