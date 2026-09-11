import io
import json

from fastapi.testclient import TestClient

from apps.api.main import create_app
from localops.config import Settings


def test_rag_and_citation(client, document):
    r = client.post(
        "/api/chat", json={"message": "What is our refund policy?", "conversation_id": "rag-test"}
    )
    assert r.status_code == 200, r.text
    data = r.json()
    assert "30 days" in data["answer"]
    assert data["citations"][0]["document_id"] == document["id"]
    assert data["trajectory"] == ["supervisor", "knowledge_agent"]
    assert len(client.get("/api/conversations/rag-test").json()) == 2


def test_insufficient_evidence(client):
    r = client.post("/api/chat", json={"message": "What is the secret code?"})
    assert "sufficient local evidence" in r.json()["answer"]
    assert not r.json()["citations"]


def test_analytics(client, dataset):
    r = client.post(
        "/api/analytics/query",
        json={
            "dataset_id": dataset["id"],
            "query": "SELECT month, SUM(revenue) AS revenue FROM dataset GROUP BY month ORDER BY month",
        },
    )
    assert r.status_code == 200, r.text
    assert r.json()["rows"] == [
        {"month": "2026-07", "revenue": 100.0},
        {"month": "2026-08", "revenue": 120.0},
    ]


def test_graph_analytics(client, dataset):
    r = client.post(
        "/api/chat", json={"message": "Analyze sales performance", "dataset_id": dataset["id"]}
    )
    assert r.status_code == 200, r.text
    assert r.json()["analytics"]["rows"][1]["revenue"] == 120


def test_graph_approve_idempotent(client):
    response = client.post(
        "/api/chat", json={"message": "Create a follow-up task", "conversation_id": "approval-test"}
    )
    assert response.status_code == 200, response.text
    data = response.json()
    approval = data["approval"]["approval_id"]
    assert client.get("/api/tasks").json() == []
    result = client.post("/api/tools/approve", json={"approval_id": approval, "approved": True})
    assert result.status_code == 200, result.text
    assert result.json()["executed"] is True
    for _ in range(3):
        client.post("/api/tools/approve", json={"approval_id": approval, "approved": True})
    assert len(client.get("/api/tasks").json()) == 1
    assert (
        len([r for r in client.get("/api/activity").json() if r["event"] == "action_approved"]) == 1
    )


def test_reject(client):
    approval = client.post("/api/chat", json={"message": "Create a task"}).json()["approval"][
        "approval_id"
    ]
    assert (
        client.post("/api/approvals/" + approval, json={"approved": False}).json()["status"]
        == "rejected"
    )
    assert client.get("/api/tasks").json() == []


def test_restart_approval(tmp_path):
    settings = Settings(data_dir=tmp_path, mode="extractive")
    with TestClient(create_app(settings)) as c:
        data = c.post(
            "/api/chat", json={"message": "Create a task", "conversation_id": "persistent"}
        ).json()
    with TestClient(create_app(settings)) as c:
        assert len(c.get("/api/conversations/persistent").json()) == 2
        r = c.post("/api/approvals/" + data["approval"]["approval_id"], json={"approved": True})
        assert r.status_code == 200, r.text
        assert len(c.get("/api/tasks").json()) == 1


def test_multistep_report(client, dataset):
    r = client.post(
        "/api/chat", json={"message": "Generate a sales report", "dataset_id": dataset["id"]}
    )
    assert r.status_code == 200, r.text
    assert r.json()["trajectory"] == ["supervisor", "analyst_agent", "action_agent"]
    key = r.json()["approval"]["approval_id"]
    client.post("/api/approvals/" + key, json={"approved": True})
    report = client.get("/api/reports/" + key)
    assert report.status_code == 200
    assert "120" in report.text


def test_delete_reindex(client, document):
    key = document["id"]
    assert client.post(f"/api/documents/{key}/reindex").status_code == 200
    assert client.delete(f"/api/documents/{key}").status_code == 200
    assert client.get(f"/api/documents/{key}").status_code == 404
    assert client.get("/api/documents/search", params={"q": "refund"}).json() == []


def test_docx_loader(client):
    from docx import Document

    doc = Document()
    doc.add_paragraph("Onboarding takes five days.")
    buffer = io.BytesIO()
    doc.save(buffer)
    r = client.post("/api/documents/upload", files={"file": ("onboarding.docx", buffer.getvalue())})
    assert r.status_code == 200, r.text


def test_xlsx_loader(client):
    from openpyxl import Workbook

    w = Workbook()
    w.active.append(["product", "stock"])
    w.active.append(["Widget", 5])
    b = io.BytesIO()
    w.save(b)
    r = client.post("/api/analytics/upload", files={"file": ("inventory.xlsx", b.getvalue())})
    assert r.status_code == 200, r.text
    assert r.json()["rows"] == [{"product": "Widget", "stock": "5"}]


def test_stream(client, document):
    with client.stream("POST", "/api/chat/stream", json={"message": "refund policy"}) as response:
        lines = list(response.iter_lines())
    assert "event: result" in lines
    payload = next(
        json.loads(lines[i + 1][6:]) for i, line in enumerate(lines) if line == "event: result"
    )
    assert "30 days" in payload["answer"]
